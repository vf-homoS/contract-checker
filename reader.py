"""
reader.py - умеет доставать текст из файла договора.
Поддерживает .txt и .docx (Word).
"""

import os


def read_contract(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".txt":
        return _read_txt(path)
    raise ValueError("Поддерживаются только файлы .txt и .docx")


def _read_txt(path: str) -> str:
    # Русские файлы бывают в разных кодировках, пробуем по очереди.
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            with open(path, "r", encoding=encoding) as file:
                return file.read()
        except UnicodeDecodeError:
            continue
    raise ValueError("Не удалось прочитать файл: непонятная кодировка")


def _read_docx(path: str) -> str:
    try:
        import docx  # библиотека python-docx
    except ImportError:
        raise RuntimeError(
            "Для чтения .docx нужна библиотека python-docx. "
            "Установи её командой: pip install python-docx")

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]

    # Реквизиты сторон часто лежат в таблицах, их тоже читаем.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(cell.text.splitlines())
    return "\n".join(parts)
